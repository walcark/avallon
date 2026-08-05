#!/usr/bin/env python3
"""Build a demo "company" reference document for the Word export.

Pandoc styles an exported note by copying the named styles of a *reference*
document: headings, body, tables, plus the custom styles the export filter
targets (colored runs, admonition boxes). This script writes such a reference,
with an invented corporate look, so the pipeline can be tested end to end. Swap
the produced file for the real company template (same style names) to get the
real branding.

    python scripts/make_reference_docx.py [-o path/reference.docx]

Style names that must exist (the export relies on them):
    Title, Heading 1..3, Normal, Rouge/Orange/Vert/Bleu/Violet (character),
    Encadré Note, Encadré Warning (paragraph).
"""

from __future__ import annotations

import argparse
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# Invented corporate palette (clearly a placeholder, not a real brand).
COMPANY = "Entreprise Modèle"
NAVY = RGBColor(0x1F, 0x4E, 0x79)
SLATE = RGBColor(0x33, 0x3A, 0x40)
COLORS = {
    "Rouge": RGBColor(0xCC, 0x22, 0x22),
    "Orange": RGBColor(0xC0, 0x6A, 0x12),
    "Vert": RGBColor(0x2E, 0x7D, 0x46),
    "Bleu": RGBColor(0x2A, 0x6D, 0xB3),
    "Violet": RGBColor(0x7A, 0x45, 0xB0),
}


def _shading(style, fill_hex: str) -> None:
    """Give *style* a solid background fill (``w:shd``)."""
    pPr = style.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    pPr.append(shd)


def _bottom_border(paragraph, color_hex: str) -> None:
    """Underline a header paragraph with a thin bottom rule."""
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color_hex)
    borders.append(bottom)
    pPr.append(borders)


def _pandoc_reference() -> Path:
    """Write Pandoc's default reference document to a temp file and return it.

    Building on top of Pandoc's own reference keeps its correct table style
    (a python-docx blank template lacks it, which collapses every exported
    table in the converters) and its full heading set; the customization below
    only recolors and adds styles.
    """
    tmp = Path(tempfile.mkdtemp()) / "pandoc-reference.docx"
    data = subprocess.run(
        ["pandoc", "--print-default-data-file", "reference.docx"],
        check=True,
        capture_output=True,
    ).stdout
    tmp.write_bytes(data)
    return tmp


def _enhance_table_style(doc) -> None:
    """Give Pandoc's ``Table`` style light gridlines and a navy header row.

    Edits the existing style in place (rather than adding a new one) so the
    layout properties Pandoc relies on stay intact. Elements are inserted at
    their schema position so the document keeps validating.
    """
    el = doc.styles["Table"].element
    tblPr = el.find(qn("w:tblPr"))

    # Light inside gridlines for readability (schema: tblBorders before
    # tblCellMar).
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "C9D3DE")
        borders.append(e)
    cell_mar = tblPr.find(qn("w:tblCellMar"))
    if cell_mar is not None:
        cell_mar.addprevious(borders)
    else:
        tblPr.append(borders)

    # Header row: a light navy-tinted fill with navy bold text. Dark-on-light
    # stays legible whatever the converter's style-precedence quirks (a table
    # style's run color can lose to the body color, so we avoid white-on-navy).
    first_row = None
    for sp in el.findall(qn("w:tblStylePr")):
        if sp.get(qn("w:type")) == "firstRow":
            first_row = sp
            break
    if first_row is None:
        first_row = OxmlElement("w:tblStylePr")
        first_row.set(qn("w:type"), "firstRow")
        el.append(first_row)

    tcPr = first_row.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        first_row.append(tcPr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "DCE6F1")
    borders_anchor = tcPr.find(qn("w:tcBorders"))
    if borders_anchor is not None:
        borders_anchor.addnext(shd)
    else:
        tcPr.insert(0, shd)

    # rPr (navy bold) comes before tcPr inside a tblStylePr.
    rPr = OxmlElement("w:rPr")
    rPr.append(OxmlElement("w:b"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    rPr.append(color)
    tcPr.addprevious(rPr)


def _para_borders(style, edges, color: str = "DDE1E6") -> None:
    """Add single borders on *edges* of a paragraph style.

    In CT_PPr, ``pBdr`` sits immediately before ``shd``, so the borders are
    anchored just before the shading (which the callers add first).
    """
    pPr = style.element.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for edge in edges:
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "6")
        e.set(qn("w:color"), color)
        borders.append(e)
    shd = pPr.find(qn("w:shd"))
    if shd is not None:
        shd.addprevious(borders)
    else:
        pPr.append(borders)


def _code_block_style(doc, normal) -> None:
    """Add the code-block card styles Pandoc's export relies on.

    ``Code Label`` is the darker header carrying the language name; ``Source
    Code`` is the lighter, monospace, syntax-highlighted body. They stack with
    no gap (the label keeps with the next paragraph, the body starts flush) and
    share a border frame, so together they read as one two-tone card. Without a
    ``Source Code`` style Pandoc invents one based on ``Normal``, which now
    inherits the justified body alignment (wrong for code).
    """
    # Header strip: dark fill, light monospace text, glued to the code below.
    label = doc.styles.add_style("Code Label", WD_STYLE_TYPE.PARAGRAPH)
    label.base_style = normal
    label.font.name = "Consolas"
    label.font.size = Pt(8)
    label.font.bold = True
    label.font.color.rgb = RGBColor(0xE8, 0xEB, 0xEE)
    lpf = label.paragraph_format
    lpf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lpf.line_spacing = 1.0
    lpf.space_before = Pt(8)
    lpf.space_after = Pt(0)
    lpf.keep_with_next = True
    _shading(label, "33383E")
    _para_borders(label, ("top", "left", "right", "bottom"))

    # Code body: light fill, monospace. The run font is set on the style so the
    # whitespace between highlighted tokens stays monospace too. No top border
    # (the label's bottom border is the divider) and no space before, so it sits
    # flush under the header.
    style = doc.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = normal
    style.font.name = "Consolas"
    style.font.size = Pt(9.5)
    style.font.color.rgb = SLATE
    pf = style.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.keep_together = True
    _shading(style, "F4F5F7")
    _para_borders(style, ("bottom", "left", "right"))


def _admonition_card(
    doc, normal, name, title_fill, title_text, body_fill, body_text, border
) -> None:
    """Add the two-paragraph card styles for an admonition *name*.

    ``Encadré <name> Titre`` is the darker header (kept with the body below);
    ``Encadré <name>`` is the lighter body. They share a border frame and stack
    flush, matching the code-block card. Consecutive body paragraphs merge into
    one frame, so the bottom border is drawn once, under the last line.
    """
    title = doc.styles.add_style(f"Encadré {name} Titre", WD_STYLE_TYPE.PARAGRAPH)
    title.base_style = normal
    title.font.bold = True
    title.font.color.rgb = title_text
    tpf = title.paragraph_format
    tpf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tpf.space_before = Pt(8)
    tpf.space_after = Pt(0)
    tpf.keep_with_next = True
    _shading(title, title_fill)
    _para_borders(title, ("top", "left", "right", "bottom"), border)

    body = doc.styles.add_style(f"Encadré {name}", WD_STYLE_TYPE.PARAGRAPH)
    body.base_style = normal
    body.font.color.rgb = body_text
    bpf = body.paragraph_format
    bpf.space_before = Pt(0)
    bpf.space_after = Pt(8)
    _shading(body, body_fill)
    _para_borders(body, ("left", "right", "bottom"), border)


def _page_number(paragraph) -> None:
    """Append a live PAGE field to *paragraph*."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def build(out_path: Path) -> None:
    """Write the reference document to *out_path*."""
    doc = Document(str(_pandoc_reference()))

    # Body text.
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = SLATE
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    # Justify body copy to match the web layout (headings keep their own
    # left-aligned styles, which do not inherit this).
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Document title and headings, in the corporate navy.
    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = NAVY

    for name, size in (("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 12)):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = NAVY
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)

    # Character styles for {couleur}(...) spans.
    for name, rgb in COLORS.items():
        style = doc.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
        style.font.color.rgb = rgb
        style.font.bold = True

    # Admonition boxes: a two-tone card, matching the code block. Dark title
    # header over a light body, tinted to each type (blue for note, amber for
    # warning).
    _admonition_card(
        doc,
        normal,
        "Note",
        title_fill="1F5A8F",
        title_text=RGBColor(0xEA, 0xF1, 0xF8),
        body_fill="EAF1F8",
        body_text=SLATE,
        border="C4DBF0",
    )
    _admonition_card(
        doc,
        normal,
        "Warning",
        title_fill="8A6A10",
        title_text=RGBColor(0xFD, 0xF3, 0xDD),
        body_fill="FDF3DD",
        body_text=RGBColor(0x7A, 0x52, 0x0A),
        border="E6D28C",
    )

    # Monospace, boxed style for exported code blocks.
    _code_block_style(doc, normal)

    # Corporate look for the table style Pandoc applies to exported tables.
    _enhance_table_style(doc)

    # Header and footer carry the company identity.
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    run = header.add_run(COMPANY)
    run.font.name = "Calibri"
    run.font.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = NAVY
    _bottom_border(header, "C9D3DE")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tag = footer.add_run(f"{COMPANY}  ·  page ")
    tag.font.size = Pt(9)
    tag.font.color.rgb = SLATE
    _page_number(footer)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Écrit : {out_path}")


def main() -> int:
    default = (
        Path(__file__).resolve().parent.parent
        / "src"
        / "pages"
        / "exporter_data"
        / "reference-entreprise.docx"
    )
    ap = argparse.ArgumentParser(description="Fabriquer un reference.docx de démo.")
    ap.add_argument("-o", "--output", type=Path, default=default)
    args = ap.parse_args()
    build(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
